#
# import os
# import tempfile
# from django.shortcuts import render
# from .forms import AudioFileForm
# from .models import AudioFile
# import speech_recognition as sr
# from docx import Document
# from django.http import HttpResponse
# from django.shortcuts import get_object_or_404
# from pydub import AudioSegment
#
#
#
# def download_word(request, audio_id):
#     audio_instance = get_object_or_404(AudioFile, pk=audio_id)
#
#     document = Document()
#     document.add_heading('متن استخراج‌شده از فایل صوتی', 0)
#     document.add_paragraph(audio_instance.transcript or "متنی وجود ندارد.")
#
#     response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
#     response['Content-Disposition'] = f'attachment; filename=transcript_{audio_id}.docx'
#
#     document.save(response)
#     return response
#
# def upload_audio(request):
#     transcript = None
#     audio_file = None
#
#     if request.method == 'POST':
#         form = AudioFileForm(request.POST, request.FILES)
#         if form.is_valid():
#             audio_file = form.save(commit=False)
#             original_file = request.FILES['audio']
#
#             # ذخیره فایل موقت اولیه
#             with tempfile.NamedTemporaryFile(delete=False, suffix=os.path.splitext(original_file.name)[1]) as temp_audio:
#                 for chunk in original_file.chunks():
#                     temp_audio.write(chunk)
#                 temp_audio.flush()
#                 temp_audio_path = temp_audio.name
#
#             file_ext = os.path.splitext(temp_audio_path)[1].lower()
#
#             if not os.path.exists(temp_audio_path):
#                 return HttpResponse("⚠ فایل صوتی موقت پیدا نشد.", status=500)
#
#             try:
#                 # تبدیل به WAV در صورت نیاز
#                 if file_ext in ['.ogg', '.webm', '.m4a']:
#                     audio = AudioSegment.from_file(temp_audio_path)
#                     temp_wav_path = temp_audio_path.replace(file_ext, '.wav')
#                     audio.export(temp_wav_path, format="wav")
#                     audio_path = temp_wav_path
#                 else:
#                     audio_path = temp_audio_path
#
#                 # تقسیم به قطعات 30 ثانیه‌ای و تبدیل به متن
#                 audio = AudioSegment.from_wav(audio_path)
#                 duration_ms = len(audio)
#                 chunk_length_ms = 30 * 1000  # ۳۰ ثانیه
#
#                 recognizer = sr.Recognizer()
#                 combined_text = ""
#
#                 for i in range(0, duration_ms, chunk_length_ms):
#                     chunk = audio[i:i + chunk_length_ms]
#
#                     # ایجاد فایل موقتی برای هر chunk و بستن آن بلافاصله
#                     temp_chunk = tempfile.NamedTemporaryFile(delete=False, suffix=".wav")
#                     temp_chunk_path = temp_chunk.name
#                     temp_chunk.close()
#
#                     chunk.export(temp_chunk_path, format="wav")
#
#                     with sr.AudioFile(temp_chunk_path) as source:
#                         audio_data = recognizer.record(source)
#
#                     try:
#                         chunk_text = recognizer.recognize_google(audio_data, language='fa-IR')
#                     except sr.UnknownValueError:
#                         chunk_text = "[نامفهوم]"
#                     except sr.RequestError:
#                         chunk_text = "[خطای ارتباط با گوگل]"
#
#                     combined_text += chunk_text + " "
#
#                     os.remove(temp_chunk_path)  # حذف فایل chunk
#
#                 transcript = combined_text.strip()
#                 audio_file.transcript = transcript
#                 audio_file.save()
#
#             finally:
#                 # پاک‌سازی فایل‌های اصلی موقت
#                 if os.path.exists(temp_audio_path):
#                     os.remove(temp_audio_path)
#                 if file_ext in ['.ogg', '.webm', '.m4a'] and os.path.exists(audio_path):
#                     os.remove(audio_path)
#
#     else:
#         form = AudioFileForm()
#
#     return render(request, 'upload.html', {
#         'form': form,
#         'transcript': transcript,
#         'audio_file': audio_file,
#     })
#

import os
import tempfile
from django.shortcuts import render, get_object_or_404
from django.http import HttpResponse
from .forms import AudioFileForm
from .models import AudioFile
try:
    from docx import Document
except ImportError:
    Document = None

try:
    from openai import OpenAI
    import openai
except ImportError:
    OpenAI = None
    openai = None



if openai:
    client = openai.OpenAI(api_key= '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'
    )
else:
    client = None

def upload_audio(request):
    transcript = None
    audio_file = None

    if request.method == 'POST':
        form = AudioFileForm(request.POST, request.FILES)
        if form.is_valid():
            audio_file = form.save(commit=False)
            original_file = request.FILES['audio']

            # ذخیره فایل موقت
            with tempfile.NamedTemporaryFile(delete=False, suffix=os.path.splitext(original_file.name)[1]) as temp_audio:
                for chunk in original_file.chunks():
                    temp_audio.write(chunk)
                temp_audio.flush()
                temp_audio_path = temp_audio.name

            try:
                # ارسال به API Whisper
                with open(temp_audio_path, "rb") as f:
                    response = client.audio.transcriptions.create(
                        model="gpt-4o-mini-transcribe",   # یا "whisper-1"
                        file=f,
                        language="fa"  # فارسی
                    )
                    transcript = response.text

                audio_file.transcript = transcript
                audio_file.save()

            finally:
                if os.path.exists(temp_audio_path):
                    os.remove(temp_audio_path)

    else:
        form = AudioFileForm()

    return render(request, 'upload.html', {
        'form': form,
        'transcript': transcript,
        'audio_file': audio_file,
    })

def download_word(request, audio_id):
    audio_instance = get_object_or_404(AudioFile, pk=audio_id)

    document = Document()
    document.add_heading('متن استخراج‌شده از فایل صوتی', 0)
    document.add_paragraph(audio_instance.transcript or "متنی وجود ندارد.")

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    response['Content-Disposition'] = f'attachment; filename=transcript_{audio_id}.docx'

    document.save(response)
    return response